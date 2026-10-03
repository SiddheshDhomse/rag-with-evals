import io
from pathlib import Path
from typing import List, Dict, Any, Union
from langchain_core.documents import Document


def load_file_to_documents(uploaded_file, filename: str) -> List[Document]:
    """
    Parses an uploaded file buffer or file path into a list of LangChain Document objects.
    Supports PDF, TXT, MD, DOCX.
    """
    ext = Path(filename).suffix.lower()
    docs = []

    if isinstance(uploaded_file, (str, Path)):
        # Local file path
        path = Path(uploaded_file)
        if ext in [".txt", ".md"]:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            docs.append(Document(page_content=content, metadata={"source": filename}))
        elif ext == ".pdf":
            import pypdf
            reader = pypdf.PdfReader(str(path))
            for page_num, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    docs.append(Document(page_content=text, metadata={"source": filename, "page": page_num + 1}))
        elif ext == ".docx":
            import docx
            doc = docx.Document(str(path))
            full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            if full_text:
                docs.append(Document(page_content=full_text, metadata={"source": filename}))
    else:
        # Streamlit UploadedFile (in-memory buffer)
        file_bytes = uploaded_file.read()
        if ext in [".txt", ".md"]:
            content = file_bytes.decode("utf-8", errors="ignore")
            docs.append(Document(page_content=content, metadata={"source": filename}))
        elif ext == ".pdf":
            import pypdf
            pdf_stream = io.BytesIO(file_bytes)
            reader = pypdf.PdfReader(pdf_stream)
            for page_num, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    docs.append(Document(page_content=text, metadata={"source": filename, "page": page_num + 1}))
        elif ext == ".docx":
            import docx
            docx_stream = io.BytesIO(file_bytes)
            doc = docx.Document(docx_stream)
            full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            if full_text:
                docs.append(Document(page_content=full_text, metadata={"source": filename}))

    return docs


def load_amnesty_qa_dataset(split: str = "eval") -> Dict[str, Any]:
    """
    Downloads and prepares the explodinggradients/amnesty_qa dataset.
    Returns:
      - 'documents': List[Document] for vector store ingestion
      - 'eval_samples': List[Dict] with questions, ground_truth, and contexts
    """
    from datasets import load_dataset
    try:
        dataset = load_dataset("explodinggradients/amnesty_qa", "english_v2", split=split)
    except Exception:
        try:
            dataset = load_dataset("explodinggradients/amnesty_qa", "english_v3", split=split)
        except Exception:
            dataset = load_dataset("explodinggradients/amnesty_qa", split=split)

    documents: List[Document] = []
    eval_samples: List[Dict[str, Any]] = []

    seen_contexts = set()

    for idx, row in enumerate(dataset):
        question = row.get("question", "")
        ground_truth = row.get("ground_truth", "")
        contexts = row.get("contexts", [])

        eval_samples.append({
            "id": idx,
            "question": question,
            "ground_truth": ground_truth,
            "reference_contexts": contexts
        })

        for c_idx, ctx in enumerate(contexts):
            if ctx not in seen_contexts:
                seen_contexts.add(ctx)
                documents.append(
                    Document(
                        page_content=ctx,
                        metadata={
                            "source": f"amnesty_report_doc_{idx+1}",
                            "context_id": f"{idx}_{c_idx}"
                        }
                    )
                )

    return {
        "documents": documents,
        "eval_samples": eval_samples,
        "total_documents": len(documents),
        "total_eval_samples": len(eval_samples)
    }
